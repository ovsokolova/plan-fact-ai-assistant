from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from docx import Document

from src.charts import make_bar, make_waterfall
from src.export import (
    export_charts_html,
    report_to_docx,
    report_to_markdown_bytes,
)
from src.llm_agent import ReportJSON, generate, report_to_markdown


@pytest.fixture
def variance_df() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "account": ["Маркетинг", "Логистика"],
            "period": ["2025-01", "2025-01"],
            "plan": [1_000_000.0, 800_000.0],
            "fact": [1_350_000.0, 760_000.0],
            "abs_variance": [350_000.0, -40_000.0],
            "rel_variance": [0.35, -0.05],
            "status": ["red", "green"],
            "owner": ["CMO", "COO"],
        }
    )


@pytest.fixture
def report(variance_df: pd.DataFrame) -> ReportJSON:
    return generate(variance_df, [])


def test_report_to_docx_returns_bytes(report: ReportJSON) -> None:
    data = report_to_docx(report)
    assert isinstance(data, bytes)
    assert len(data) > 1000


def test_report_to_docx_readable(report: ReportJSON) -> None:
    data = report_to_docx(report)
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Аналитическая записка" in text
    assert "Маркетинг" in text
    assert "Резюме" in text


def test_report_to_docx_contains_hypotheses(report: ReportJSON) -> None:
    data = report_to_docx(report)
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Гипотезы причин" in text
    assert "Вопросы ответственным" in text


def test_report_to_markdown_bytes(report: ReportJSON) -> None:
    md = report_to_markdown(report)
    data = report_to_markdown_bytes(md)
    assert isinstance(data, bytes)
    assert "Аналитическая записка" in data.decode("utf-8")


def test_export_charts_html_creates_files(tmp_path: Path, variance_df: pd.DataFrame) -> None:
    charts = {
        "bar": make_bar(variance_df),
        "waterfall": make_waterfall(variance_df),
    }
    files = export_charts_html(charts, output_dir=tmp_path / "charts")
    assert len(files) == 2
    for f in files:
        assert f.exists()
        assert f.suffix == ".html"
        assert f.stat().st_size > 500


def test_export_charts_html_sanitizes_names(tmp_path: Path, variance_df: pd.DataFrame) -> None:
    charts = {"plan vs fact / 2025": make_bar(variance_df)}
    files = export_charts_html(charts, output_dir=tmp_path / "charts")
    assert len(files) == 1
    # недопустимые символы заменены на _
    assert "/" not in files[0].name
    assert " " not in files[0].name


def test_export_charts_html_empty(tmp_path: Path) -> None:
    files = export_charts_html({}, output_dir=tmp_path / "empty")
    assert files == []
    assert (tmp_path / "empty").exists()

def test_report_to_pdf_does_not_crash(report: ReportJSON) -> None:
    """PDF-экспорт не должен падать, даже если WeasyPrint не установлен."""
    from src.export import report_to_pdf

    result = report_to_pdf(report)
    # либо bytes (WeasyPrint установлен), либо None (нет WeasyPrint)
    assert result is None or isinstance(result, bytes)
    if isinstance(result, bytes):
        assert len(result) > 500


def test_markdown_to_html_simple() -> None:
    from src.export import _markdown_to_simple_html

    md = "# Заголовок\n\n- пункт 1\n- пункт 2\n\n**жирный** текст"
    html = _markdown_to_simple_html(md)
    assert "<h1>Заголовок</h1>" in html
    assert "<li>пункт 1</li>" in html
    assert "<strong>жирный</strong>" in html
